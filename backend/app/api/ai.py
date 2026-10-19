from io import BytesIO
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException
from docx import Document as WordDocument
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from api.dashboard import current_user
from database import get_db
from models.ai_feedback import AIFeedback
from models.document import Document
from models.user import User
from schemas.documents import (
    FeedbackRequest,
    FeedbackResponse,
    ApplyChangeRequest,
    ApplyChangeResponse,
    PredictionRequest,
    PredictionResponse,
)
from schemas.structured_documents import (
    DocumentGenerationRequest,
    DocumentGenerationResponse,
)
from services.ai_service import AIService


router = APIRouter(prefix="/ai", tags=["ai"])


# ---------------------------------------------------------
# Helpers
# ---------------------------------------------------------

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".json",
    ".xml",
    ".html",
    ".htm",
}


def _iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                yield paragraph

            for nested_table in cell.tables:
                yield from _iter_table_paragraphs(nested_table)


def _iter_document_paragraphs(document):
    """
    Walk through:
    - normal paragraphs
    - tables
    - headers
    - footers
    """

    for paragraph in document.paragraphs:
        yield paragraph

    for table in document.tables:
        yield from _iter_table_paragraphs(table)

    for section in document.sections:
        parts = (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        )

        for part in parts:
            for paragraph in part.paragraphs:
                yield paragraph

            for table in part.tables:
                yield from _iter_table_paragraphs(table)


def _replace_first_in_paragraph(
    paragraph,
    matched_text: str,
    proposed_text: str,
) -> bool:
    """
    Replace text even when matched_text is split across
    multiple Word runs.

    Example:

        run 1 = "Hello "
        run 2 = "world"

    matched_text = "Hello world"

    Normal run.text.replace() would fail.
    """

    runs = paragraph.runs

    if not runs:
        return False

    run_texts = [run.text or "" for run in runs]
    full_text = "".join(run_texts)

    start = full_text.find(matched_text)

    if start == -1:
        return False

    end = start + len(matched_text)

    # Find start run
    cursor = 0
    start_run_index = None

    for index, text in enumerate(run_texts):
        next_cursor = cursor + len(text)

        if cursor <= start < next_cursor:
            start_run_index = index
            break

        cursor = next_cursor

    # Empty matched text or weird boundary case
    if start_run_index is None:
        return False

    # Find end run
    cursor = 0
    end_run_index = None

    for index, text in enumerate(run_texts):
        next_cursor = cursor + len(text)

        if cursor < end <= next_cursor:
            end_run_index = index
            break

        cursor = next_cursor

    if end_run_index is None:
        return False

    # Position inside starting run
    start_run_offset = start - sum(
        len(run_texts[i]) for i in range(start_run_index)
    )

    # Position inside ending run
    end_run_offset = end - sum(
        len(run_texts[i]) for i in range(end_run_index)
    )

    if start_run_index == end_run_index:
        original = run_texts[start_run_index]

        runs[start_run_index].text = (
            original[:start_run_offset]
            + proposed_text
            + original[end_run_offset:]
        )

        return True

    # Text before matched_text
    prefix = run_texts[start_run_index][:start_run_offset]

    # Text after matched_text
    suffix = run_texts[end_run_index][end_run_offset:]

    # Put replacement into first run
    runs[start_run_index].text = prefix + proposed_text

    # Clear runs in between
    for index in range(start_run_index + 1, end_run_index):
        runs[index].text = ""

    # Preserve text after the match
    runs[end_run_index].text = suffix

    return True


def _replace_first_in_docx(
    source_data: bytes,
    matched_text: str,
    proposed_text: str,
) -> bytes | None:
    """
    Modify DOCX in memory.

    Returns:
        edited bytes
        None if matched_text was not found
    """

    document = WordDocument(BytesIO(source_data))

    for paragraph in _iter_document_paragraphs(document):
        changed = _replace_first_in_paragraph(
            paragraph,
            matched_text,
            proposed_text,
        )

        if changed:
            output = BytesIO()
            document.save(output)
            return output.getvalue()

    return None


def _replace_first_in_text_file(
    source_data: bytes,
    matched_text: str,
    proposed_text: str,
) -> bytes | None:
    try:
        text = source_data.decode("utf-8")
    except UnicodeDecodeError:
        raise HTTPException(
            status_code=422,
            detail="This text file is not valid UTF-8.",
        )

    if matched_text not in text:
        return None

    updated = text.replace(
        matched_text,
        proposed_text,
        1,
    )

    return updated.encode("utf-8")


# ---------------------------------------------------------
# Generate
# ---------------------------------------------------------

@router.post(
    "/generate",
    response_model=DocumentGenerationResponse,
)
def generate_document(
    payload: DocumentGenerationRequest,
    _user: User = Depends(current_user),
) -> DocumentGenerationResponse:
    """
    Generate a document proposal.

    The AI provider does not receive the user's auth token.
    The proposal is not automatically persisted.
    """

    if not payload.prompt or not payload.prompt.strip():
        raise HTTPException(
            status_code=422,
            detail="Prompt cannot be empty.",
        )

    try:
        return AIService.generate_document(
            payload.prompt.strip()
        )
    except Exception:
        raise HTTPException(
            status_code=502,
            detail="AI document generation failed.",
        )


# ---------------------------------------------------------
# Predict
# ---------------------------------------------------------

@router.post(
    "/predict",
    response_model=PredictionResponse,
)
def predict(
    payload: PredictionRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> PredictionResponse:

    document = db.get(
        Document,
        payload.document_id,
    )

    if document is None:
        raise HTTPException(
            status_code=404,
            detail="Document not found.",
        )

    if document.owner_id != user.id:
        raise HTTPException(
            status_code=404,
            detail="Document not found.",
        )

    if not payload.command or not payload.command.strip():
        raise HTTPException(
            status_code=422,
            detail="Command cannot be empty.",
        )

    extracted_text = document.extracted_text or ""

    try:
        prediction = AIService.predict(
            payload.command.strip(),
            extracted_text,
        )
    except Exception:
        raise HTTPException(
            status_code=502,
            detail="AI prediction failed.",
        )

    return PredictionResponse(
        document_id=document.id,
        intent=prediction.intent,
        confidence=prediction.confidence,
        matched_text=prediction.matched_text,
        proposed_text=prediction.proposed_text,
        explanation=prediction.explanation,
    )


# ---------------------------------------------------------
# Feedback
# ---------------------------------------------------------

@router.post(
    "/feedback",
    response_model=FeedbackResponse,
    status_code=201,
)
def feedback(
    payload: FeedbackRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> FeedbackResponse:

    if payload.document_id is not None:

        document = db.get(
            Document,
            payload.document_id,
        )

        if document is None or document.owner_id != user.id:
            raise HTTPException(
                status_code=404,
                detail="Document not found.",
            )

    record = AIFeedback(
        user_id=user.id,
        document_id=payload.document_id,
        command=payload.command,
        prediction=payload.prediction,
        accepted=payload.accepted,
        correction=payload.correction,
    )

    try:
        db.add(record)
        db.commit()
        db.refresh(record)
    except SQLAlchemyError:
        db.rollback()

        raise HTTPException(
            status_code=500,
            detail="Failed to save AI feedback.",
        )

    return FeedbackResponse(
        id=record.id,
        message="Feedback recorded for future model improvement.",
    )


# ---------------------------------------------------------
# Apply AI change
# ---------------------------------------------------------

@router.post(
    "/apply",
    response_model=ApplyChangeResponse,
)
def apply_change(
    payload: ApplyChangeRequest,
    document_id: int,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> ApplyChangeResponse:

    document = db.get(
        Document,
        document_id,
    )

    if document is None or document.owner_id != user.id:
        raise HTTPException(
            status_code=404,
            detail="Document not found.",
        )

    # Structured documents must use their own schema/update API.
    if (
        document.content_type
        == "application/vnd.seqotia.structured-document+json"
    ):
        raise HTTPException(
            status_code=422,
            detail=(
                "Update structured documents through "
                "their validated field schema."
            ),
        )

    matched_text = payload.matched_text
    proposed_text = payload.proposed_text

    if not matched_text:
        raise HTTPException(
            status_code=422,
            detail="matched_text cannot be empty.",
        )

    if proposed_text is None:
        raise HTTPException(
            status_code=422,
            detail="proposed_text cannot be null.",
        )

    original_extracted_text = document.extracted_text or ""

    # First make sure AI prediction is still applicable.
    if matched_text not in original_extracted_text:
        return ApplyChangeResponse(
            document_id=document.id,
            changed=False,
            message=(
                "The original text was not found. "
                "Nothing was changed."
            ),
        )

    extension = Path(
        document.filename or ""
    ).suffix.lower()

    edited_data = None

    # -----------------------------------------------------
    # DOCX
    # -----------------------------------------------------

    if extension == ".docx":

        if not document.source_data:
            raise HTTPException(
                status_code=422,
                detail="Original DOCX data is missing.",
            )

        edited_data = _replace_first_in_docx(
            document.source_data,
            matched_text,
            proposed_text,
        )

        if edited_data is None:
            return ApplyChangeResponse(
                document_id=document.id,
                changed=False,
                message=(
                    "The text was found in extracted content, "
                    "but could not be located inside the DOCX. "
                    "Nothing was changed."
                ),
            )

    # -----------------------------------------------------
    # Plain text files
    # -----------------------------------------------------

    elif extension in TEXT_EXTENSIONS:

        if not document.source_data:
            raise HTTPException(
                status_code=422,
                detail="Original file data is missing.",
            )

        edited_data = _replace_first_in_text_file(
            document.source_data,
            matched_text,
            proposed_text,
        )

        if edited_data is None:
            return ApplyChangeResponse(
                document_id=document.id,
                changed=False,
                message=(
                    "The original text was not found. "
                    "Nothing was changed."
                ),
            )

    # -----------------------------------------------------
    # Unsupported binary files
    # -----------------------------------------------------

    else:
        raise HTTPException(
            status_code=422,
            detail=(
                f"Automatic editing is not supported for "
                f"'{extension or 'this file type'}' yet."
            ),
        )

    # -----------------------------------------------------
    # Update database only after actual file modification
    # succeeded.
    # -----------------------------------------------------

    new_extracted_text = original_extracted_text.replace(
        matched_text,
        proposed_text,
        1,
    )

    document.extracted_text = new_extracted_text
    document.edited_data = edited_data

    try:
        db.commit()

    except SQLAlchemyError:
        db.rollback()

        raise HTTPException(
            status_code=500,
            detail="Failed to save the document change.",
        )

    return ApplyChangeResponse(
        document_id=document.id,
        changed=True,
        message=(
            "AI change applied successfully. "
            "The original upload remains preserved."
        ),
    )