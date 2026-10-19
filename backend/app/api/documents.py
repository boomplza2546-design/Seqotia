
from __future__ import annotations

import logging
import os
import re
from io import BytesIO
from pathlib import Path

from fastapi import (
    APIRouter,
    Depends,
    File,
    HTTPException,
    Query,
    UploadFile,
    status,
)
from fastapi.responses import StreamingResponse
from PIL import Image, ImageDraw, ImageFont, UnidentifiedImageError
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from sqlalchemy import delete, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session
from docx import Document as WordDocument

try:
    from pypdf import PdfReader
except ImportError:  # pragma: no cover - optional dependency
    PdfReader = None  # type: ignore[assignment]

try:
    import pytesseract
except ImportError:  # pragma: no cover - optional dependency
    pytesseract = None  # type: ignore[assignment]


from api.dependencies.auth import current_user
from database import get_db
from models.document import Document
from models.document_structure import DocumentStructure
from models.user import User
from schemas.documents import DocumentPublic
from schemas.structured_documents import (
    DocumentType,
    StructuredDocumentCreate,
    StructuredDocumentPublic,
    StructuredDocumentUpdate,
    TemplatePublic,
)
from services.document_rules import DocumentRules
from services.template_engine import TemplateEngine


logger = logging.getLogger(__name__)

router = APIRouter(
    prefix="/documents",
    tags=["documents"],
)


# ============================================================================
# Configuration
# ============================================================================

SUPPORTED_EXTENSIONS = {
    "txt",
    "md",
    "markdown",
    "docx",
    "pdf",
    "jpg",
    "jpeg",
    "png",
}

MAX_DOCUMENT_SIZE_MB = int(
    os.getenv("MAX_DOCUMENT_SIZE_MB", "25")
)

MAX_DOCUMENT_SIZE = (
    MAX_DOCUMENT_SIZE_MB * 1024 * 1024
)


STRUCTURED_CONTENT_TYPE = (
    "application/vnd.seqotia.structured-document+json"
)

DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "wordprocessingml.document"
)


# ============================================================================
# Filename helpers
# ============================================================================

_INVALID_FILENAME_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1F]+')


def _safe_filename(filename: str | None) -> str:
    """
    Sanitize uploaded filename.

    Never trust a client-provided filename.
    """

    value = (
        filename.strip()
        if filename
        else "untitled.txt"
    )

    value = Path(value).name
    value = _INVALID_FILENAME_CHARS.sub("_", value)

    if not value or value in {".", ".."}:
        value = "untitled.txt"

    return value[:255]


def _filename_stem(filename: str) -> str:
    return Path(filename).stem or "document"


# ============================================================================
# Document ownership
# ============================================================================

def _owned_document(
    document_id: int,
    user: User,
    db: Session,
) -> Document:
    document = db.get(
        Document,
        document_id,
    )

    if document is None or document.owner_id != user.id:
        # Do not reveal whether a document exists
        # but belongs to somebody else.
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Document not found.",
        )

    return document


# ============================================================================
# DOCX extraction
# ============================================================================

def _iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                yield paragraph

            for nested_table in cell.tables:
                yield from _iter_table_paragraphs(nested_table)


def _extract_docx_text(document: WordDocument) -> str:
    """
    Extract text from:
    - normal paragraphs
    - tables
    - headers
    - footers
    """

    chunks: list[str] = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            chunks.append(paragraph.text)

    for table in document.tables:
        for paragraph in _iter_table_paragraphs(table):
            if paragraph.text.strip():
                chunks.append(paragraph.text)

    for section in document.sections:
        parts = (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        )

        for part in parts:
            for paragraph in part.paragraphs:
                if paragraph.text.strip():
                    chunks.append(paragraph.text)

            for table in part.tables:
                for paragraph in _iter_table_paragraphs(table):
                    if paragraph.text.strip():
                        chunks.append(paragraph.text)

    return "\n".join(chunks).strip()


def _validate_docx(source_data: bytes) -> None:
    """
    Open DOCX to make sure it is actually readable as a Word file.
    """

    try:
        document = WordDocument(
            BytesIO(source_data)
        )
        _ = document.core_properties
    except Exception as error:
        logger.warning(
            "Invalid DOCX upload: %s",
            error,
        )
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded DOCX file is invalid or corrupted.",
        ) from error


# ============================================================================
# PDF extraction
# ============================================================================

def _extract_pdf_text(source_data: bytes) -> str:
    """
    Real PDF text extraction using pypdf.

    Scanned/image-only PDFs still require OCR.
    """

    if PdfReader is None:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail=(
                "PDF text extraction is unavailable. "
                "Install the 'pypdf' package."
            ),
        )

    try:
        reader = PdfReader(
            BytesIO(source_data)
        )
    except Exception as error:
        logger.warning(
            "Invalid PDF upload: %s",
            error,
        )
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded PDF file is invalid or corrupted.",
        ) from error

    chunks: list[str] = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):
        try:
            text = page.extract_text() or ""
        except Exception as error:
            logger.warning(
                "PDF page extraction failed: page=%s error=%s",
                page_number,
                error,
            )
            text = ""

        if text.strip():
            chunks.append(
                f"[Page {page_number}]\n{text.strip()}"
            )

    return "\n\n".join(chunks).strip()


# ============================================================================
# Image extraction / OCR
# ============================================================================

def _extract_image_text(
    source_data: bytes,
) -> tuple[str, str]:
    """
    Validate image and optionally perform OCR.

    Returns:
        (content, content_type)
    """

    try:
        image = Image.open(
            BytesIO(source_data)
        )

        # Force decode so malformed images are detected.
        image.load()

    except UnidentifiedImageError as error:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded image is invalid.",
        ) from error

    except Exception as error:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded image could not be read.",
        ) from error

    if pytesseract is not None:
        try:
            ocr_text = pytesseract.image_to_string(
                image
            ).strip()

            if ocr_text:
                return (
                    ocr_text,
                    f"image/{image.format.lower()}"
                    if image.format
                    else "image/png",
                )

        except Exception as error:
            logger.warning(
                "OCR failed: %s",
                error,
            )

    width, height = image.size

    return (
        (
            f"Image document imported "
            f"({width}x{height}). "
            "OCR text was not available."
        ),
        (
            f"image/{image.format.lower()}"
            if image.format
            else "image/png"
        ),
    )


# ============================================================================
# Plain text extraction
# ============================================================================

def _extract_text_file(
    source_data: bytes,
) -> str:
    try:
        return source_data.decode(
            "utf-8-sig"
        ).strip()
    except UnicodeDecodeError as error:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail=(
                "Text files must be encoded as UTF-8."
            ),
        ) from error


# ============================================================================
# Content validation
# ============================================================================

def _validate_magic_bytes(
    extension: str,
    source_data: bytes,
) -> None:
    """
    Do not trust file.content_type alone.
    """

    if extension == "pdf":
        if not source_data.startswith(b"%PDF"):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The file is not a valid PDF.",
            )

    elif extension == "docx":
        # DOCX is a ZIP/OpenXML container.
        if not source_data.startswith(b"PK"):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The file is not a valid DOCX.",
            )

    elif extension in {"jpg", "jpeg", "png"}:
        try:
            image = Image.open(
                BytesIO(source_data)
            )
            image.verify()
        except Exception as error:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The uploaded image is invalid.",
            ) from error


# ============================================================================
# Structured document response
# ============================================================================

def _structured_response(
    document: Document,
    structure: DocumentStructure,
) -> StructuredDocumentPublic:

    data = StructuredDocumentUpdate.model_validate(
        {
            "template_id": structure.template_id,
            "data": structure.data,
        }
    ).data

    return StructuredDocumentPublic(
        document_id=document.id,
        filename=document.filename,
        status=document.status,
        template_id=structure.template_id,
        version=structure.version,
        data=data,
        validation=DocumentRules.validate(data),
        created_at=structure.created_at,
        updated_at=structure.updated_at,
    )


# ============================================================================
# Document types
# ============================================================================

@router.get(
    "/types",
    response_model=list[DocumentType],
)
def list_document_types() -> list[DocumentType]:
    return list(DocumentType)


@router.get(
    "/templates",
    response_model=list[TemplatePublic],
)
def list_templates() -> list[TemplatePublic]:
    return TemplateEngine.list_templates()


# ============================================================================
# Structured documents
# ============================================================================

@router.post(
    "/structured",
    response_model=StructuredDocumentPublic,
    status_code=status.HTTP_201_CREATED,
)
def create_structured_document(
    payload: StructuredDocumentCreate,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> StructuredDocumentPublic:

    try:
        rendered = TemplateEngine.render(
            payload.template_id,
            payload.data,
        )
    except ValueError as error:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(error),
        ) from error

    validation = DocumentRules.validate(
        payload.data
    )

    filename = _safe_filename(
        payload.filename
    )

    if "." not in filename:
        filename = f"{filename}.docx"

    document = Document(
        owner_id=user.id,
        filename=filename,
        content_type=STRUCTURED_CONTENT_TYPE,
        extracted_text=rendered,
        source_data=rendered.encode("utf-8"),
        status=(
            "ready"
            if validation.valid
            else "draft"
        ),
    )

    structure: DocumentStructure | None = None

    try:
        db.add(document)
        db.flush()

        structure = DocumentStructure(
            document_id=document.id,
            document_type=(
                payload.data.document_type.value
            ),
            template_id=payload.template_id,
            data=payload.data.model_dump(
                mode="json"
            ),
        )

        db.add(structure)
        db.commit()

        db.refresh(document)
        db.refresh(structure)

    except SQLAlchemyError as error:
        db.rollback()

        logger.exception(
            "Failed to create structured document."
        )

        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Failed to create document.",
        ) from error

    return _structured_response(
        document,
        structure,
    )


# ============================================================================
# Upload
# ============================================================================

@router.post(
    "",
    response_model=DocumentPublic,
    status_code=status.HTTP_201_CREATED,
)
async def upload_document(
    file: UploadFile = File(...),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> Document:

    filename = _safe_filename(
        file.filename
    )

    extension = (
        filename.rsplit(".", 1)[-1].lower()
        if "." in filename
        else "txt"
    )

    if extension not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail=(
                "Seqotia currently supports "
                "TXT, Markdown, DOCX, PDF, JPG, "
                "and PNG files."
            ),
        )

    source_data = await file.read(
        MAX_DOCUMENT_SIZE + 1
    )

    if len(source_data) > MAX_DOCUMENT_SIZE:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=(
                f"Maximum document size is "
                f"{MAX_DOCUMENT_SIZE_MB} MB."
            ),
        )

    if not source_data:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded document is empty.",
        )

    _validate_magic_bytes(
        extension,
        source_data,
    )

    # ----------------------------------------------------
    # Extract content
    # ----------------------------------------------------

    if extension == "docx":

        _validate_docx(
            source_data
        )

        word_document = WordDocument(
            BytesIO(source_data)
        )

        content = _extract_docx_text(
            word_document
        )

        content_type = DOCX_CONTENT_TYPE

    elif extension == "pdf":

        content = _extract_pdf_text(
            source_data
        )

        content_type = (
            file.content_type
            or "application/pdf"
        )

    elif extension in {
        "jpg",
        "jpeg",
        "png",
    }:

        content, detected_content_type = (
            _extract_image_text(
                source_data
            )
        )

        content_type = detected_content_type

    else:

        content = _extract_text_file(
            source_data
        )

        content_type = (
            file.content_type
            or "text/plain"
        )

    if not content.strip():
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "No readable text was found in "
                "the document."
            ),
        )

    document = Document(
        owner_id=user.id,
        filename=filename,
        content_type=content_type,
        extracted_text=content,
        source_data=source_data,
        status="ready",
    )

    try:
        db.add(document)
        db.commit()
        db.refresh(document)

    except SQLAlchemyError as error:
        db.rollback()

        logger.exception(
            "Failed to store document."
        )

        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Failed to save document.",
        ) from error

    return document


# ============================================================================
# List documents
# ============================================================================

@router.get(
    "",
    response_model=list[DocumentPublic],
)
def list_documents(
    response,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> list[Document]:

    response.headers["Cache-Control"] = "no-store"

    return list(
        db.scalars(
            select(Document)
            .where(
                Document.owner_id == user.id
            )
            .order_by(
                Document.created_at.desc()
            )
        )
    )


# ============================================================================
# Get a single document
# ============================================================================

@router.get(
    "/{document_id}",
    response_model=DocumentPublic,
)
def get_document(
    document_id: int,
    response,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> Document:

    document = _owned_document(
        document_id,
        user,
        db,
    )

    response.headers["Cache-Control"] = "no-store"

    return document


# ============================================================================
# Delete document
# ============================================================================

@router.delete(
    "/{document_id}",
    status_code=status.HTTP_204_NO_CONTENT,
)
def delete_document(
    document_id: int,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> None:

    document = _owned_document(
        document_id,
        user,
        db,
    )

    try:
        # Explicitly remove structure first.
        db.execute(
            delete(DocumentStructure).where(
                DocumentStructure.document_id
                == document.id
            )
        )

        db.delete(document)
        db.commit()

    except SQLAlchemyError as error:
        db.rollback()

        logger.exception(
            "Failed to delete document id=%s",
            document_id,
        )

        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Failed to delete document.",
        ) from error


# ============================================================================
# Structured document
# ============================================================================

@router.get(
    "/{document_id}/structure",
    response_model=StructuredDocumentPublic,
)
def get_document_structure(
    document_id: int,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> StructuredDocumentPublic:

    document = _owned_document(
        document_id,
        user,
        db,
    )

    structure = db.scalar(
        select(DocumentStructure).where(
            DocumentStructure.document_id
            == document.id
        )
    )

    if structure is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=(
                "This upload has not been converted "
                "to a structured document yet."
            ),
        )

    return _structured_response(
        document,
        structure,
    )


@router.put(
    "/{document_id}/structure",
    response_model=StructuredDocumentPublic,
)
def update_document_structure(
    document_id: int,
    payload: StructuredDocumentUpdate,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> StructuredDocumentPublic:

    document = _owned_document(
        document_id,
        user,
        db,
    )

    structure = db.scalar(
        select(DocumentStructure).where(
            DocumentStructure.document_id
            == document.id
        )
    )

    if structure is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=(
                "This upload has not been converted "
                "to a structured document yet."
            ),
        )

    template_id = (
        payload.template_id
        or structure.template_id
    )

    try:
        rendered = TemplateEngine.render(
            template_id,
            payload.data,
        )
    except ValueError as error:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(error),
        ) from error

    validation = DocumentRules.validate(
        payload.data
    )

    structure.template_id = template_id
    structure.document_type = (
        payload.data.document_type.value
    )
    structure.data = payload.data.model_dump(
        mode="json"
    )
    structure.version += 1

    document.extracted_text = rendered
    document.status = (
        "ready"
        if validation.valid
        else "draft"
    )

    try:
        db.commit()

        db.refresh(document)
        db.refresh(structure)

    except SQLAlchemyError as error:
        db.rollback()

        logger.exception(
            "Failed to update structured document id=%s",
            document_id,
        )

        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Failed to update document.",
        ) from error

    return _structured_response(
        document,
        structure,
    )


# ============================================================================
# Validation
# ============================================================================

@router.get(
    "/{document_id}/validation",
)
def validate_document(
    document_id: int,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):

    document = _owned_document(
        document_id,
        user,
        db,
    )

    structure = db.scalar(
        select(DocumentStructure).where(
            DocumentStructure.document_id
            == document.id
        )
    )

    if structure is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=(
                "This upload has not been converted "
                "to a structured document yet."
            ),
        )

    data = StructuredDocumentUpdate.model_validate(
        {
            "template_id": structure.template_id,
            "data": structure.data,
        }
    ).data

    return DocumentRules.validate(data)


# ============================================================================
# Exporters
# ============================================================================

def _export_docx(
    document: Document,
) -> BytesIO:

    if (
        document.content_type
        == STRUCTURED_CONTENT_TYPE
    ):
        word_document = WordDocument()

        lines = (
            document.extracted_text.splitlines()
            or [""]
        )

        for line in lines:
            word_document.add_paragraph(
                line
            )

        output = BytesIO()

        word_document.save(
            output
        )

        output.seek(0)
        return output

    source = (
        document.edited_data
        or document.source_data
    )

    if not source:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Document data is unavailable.",
        )

    output = BytesIO(
        source
    )

    output.seek(0)

    return output


def _export_pdf(
    document: Document,
) -> BytesIO:

    output = BytesIO()

    pdf = canvas.Canvas(
        output,
        pagesize=A4,
    )

    _, height = A4

    y = height - 48

    pdf.setFont(
        "Helvetica",
        11,
    )

    for line in (
        document.extracted_text.splitlines()
        or [""]
    ):
        if y < 48:
            pdf.showPage()
            pdf.setFont(
                "Helvetica",
                11,
            )
            y = height - 48

        pdf.drawString(
            48,
            y,
            line[:130],
        )

        y -= 16

    pdf.save()

    output.seek(0)

    return output


def _export_image(
    document: Document,
    extension: str,
) -> BytesIO:

    lines = (
        document.extracted_text.splitlines()
        or [""]
    )

    image = Image.new(
        "RGB",
        (
            1600,
            max(
                900,
                100 + len(lines) * 28,
            ),
        ),
        "white",
    )

    draw = ImageDraw.Draw(
        image
    )

    font = ImageFont.load_default(
        size=22
    )

    for index, line in enumerate(lines):
        draw.text(
            (
                70,
                60 + index * 28,
            ),
            line[:120],
            fill="#18231d",
            font=font,
        )

    output = BytesIO()

    image.save(
        output,
        format=(
            "PNG"
            if extension == "png"
            else "JPEG"
        ),
        quality=95,
    )

    output.seek(0)

    return output


# ============================================================================
# Export
# ============================================================================

@router.get(
    "/{document_id}/export",
)
def export_document(
    document_id: int,
    format: str = Query(
        default="docx",
        pattern="^(docx|pdf|png|jpg|jpeg)$",
    ),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
    response: Response | None = None,
) -> StreamingResponse:

    if response is not None:
        response.headers["Cache-Control"] = "no-store"
        response.headers["Pragma"] = "no-cache"

    document = _owned_document(
        document_id,
        user,
        db,
    )

    structure = db.scalar(
        select(DocumentStructure).where(
            DocumentStructure.document_id
            == document.id
        )
    )

    # Structured documents must be valid before export.
    if structure is not None:

        data = StructuredDocumentUpdate.model_validate(
            {
                "template_id": structure.template_id,
                "data": structure.data,
            }
        ).data

        validation = DocumentRules.validate(
            data
        )

        if not validation.valid:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                detail={
                    "message": (
                        "Resolve validation errors "
                        "before export."
                    ),
                    "issues": [
                        issue.model_dump()
                        for issue in validation.issues
                    ],
                },
            )

    export_format = format.lower()

    exporters = {
        "docx": (
            _export_docx,
            DOCX_CONTENT_TYPE,
        ),
        "pdf": (
            _export_pdf,
            "application/pdf",
        ),
        "png": (
            lambda item: _export_image(
                item,
                "png",
            ),
            "image/png",
        ),
        "jpg": (
            lambda item: _export_image(
                item,
                "jpg",
            ),
            "image/jpeg",
        ),
        "jpeg": (
            lambda item: _export_image(
                item,
                "jpg",
            ),
            "image/jpeg",
        ),
    }

    exporter, media_type = exporters[
        export_format
    ]

    filename = _safe_filename(
        document.filename
    )

    stem = _filename_stem(
        filename
    )

    output_filename = (
        f"{stem}.{export_format}"
    )

    response.headers["Cache-Control"] = (
        "no-store"
    )

    return StreamingResponse(
        exporter(document),
        media_type=media_type,
        headers={
            "Content-Disposition": (
                f'attachment; filename="{output_filename}"'
            ),
            "X-Content-Type-Options": "nosniff",
        },
    )

